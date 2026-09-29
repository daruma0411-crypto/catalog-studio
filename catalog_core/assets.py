import csv
import hashlib
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from .importer import checked_zip,save_image,xml


def spreadsheet_preview(data,filename):
    suffix=Path(filename).suffix.lower()
    sheets=[]
    if suffix in ('.csv','.tsv'):
        try: decoded=data.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: decoded=data.decode('cp932')
            except UnicodeDecodeError as e: raise ValueError('CSVの文字コードはUTF-8またはShift-JISにしてください。') from e
        rows=[]
        for i,row in enumerate(csv.reader(io.StringIO(decoded),delimiter='\t' if suffix=='.tsv' else ',')):
            if i==201: break
            rows.append([str(v)[:2000] for v in row[:40]])
        sheets=[{'name':Path(filename).name,'rows':rows[:200],'truncated':len(rows)>200}]
    elif suffix=='.xlsx':
        with checked_zip(data) as z:
            for n in z.namelist():
                if n.endswith('.xml'): xml(z.read(n))
        from openpyxl import load_workbook
        try:
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            for ws in wb.worksheets[:10]:
                rows=[]
                for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row or 200,200),max_col=min(ws.max_column or 40,40),values_only=True):
                    rows.append(['' if v is None else str(v)[:2000] for v in row])
                sheets.append({'name':ws.title,'rows':rows,'truncated':(ws.max_row or 0)>200 or (ws.max_column or 0)>40})
            wb.close()
        except (ValueError,KeyError,OSError) as e:
            raise ValueError('Excelファイルを読み取れません。xlsx形式で保存し直してください。') from e
    elif suffix=='.xls':
        raise ValueError('旧形式.xlsは未対応です。Excelで.xlsxまたは.csvに保存して取り込んでください。')
    else:
        raise ValueError('対応する表形式は.xlsx / .csv / .tsvです。')
    return {'sheets':sheets,'note':'先頭200行・40列・10シートまでを表示。数式は文字列のまま保持し実行しません。'}


def ingest_attachment(data,filename,assets_dir):
    if not isinstance(filename,str) or not filename or len(filename)>240: raise ValueError('ファイル名が不正です。')
    if len(data)>30*1024*1024: raise ValueError('素材は1ファイル30MB以下にしてください。')
    filename=filename.replace('\\','/').split('/')[-1]
    suffix=Path(filename).suffix.lower()
    allowed={'.png','.jpg','.jpeg','.webp','.psd','.tif','.tiff','.xlsx','.csv','.tsv','.pdf','.docx','.txt'}
    if suffix not in allowed:
        if suffix=='.xls': raise ValueError('旧形式.xlsは.xlsxへ保存し直してください。')
        raise ValueError('このファイル形式には対応していません。')
    directory=Path(assets_dir)
    aid=hashlib.sha256(data).hexdigest()[:24]+suffix
    attachment={'id':aid,'name':filename,'size':len(data),'kind':'file'}
    assets={}
    if suffix in {'.png','.jpg','.jpeg','.webp','.psd','.tif','.tiff'}:
        preview=save_image(data,directory)
        assets[preview]={'id':preview,'name':filename,'kind':'image'}
        attachment.update(kind='image',preview_asset=preview)
    elif suffix in {'.xlsx','.csv','.tsv'}:
        attachment.update(kind='spreadsheet',preview=spreadsheet_preview(data,filename))
    elif suffix=='.docx':
        with checked_zip(data) as z:
            if 'word/document.xml' not in z.namelist(): raise ValueError('Wordファイルとして読み取れません。')
            root=xml(z.read('word/document.xml'))
            content='\n'.join(''.join(e.text or '' for e in p.iter() if e.tag.endswith('}t')) for p in root.iter() if p.tag.endswith('}p'))
            attachment['text_preview']=content[:30000]
    elif suffix=='.txt':
        attachment['text_preview']=data.decode('utf-8-sig',errors='replace')[:30000]
    elif suffix=='.pdf' and not data.startswith(b'%PDF-'):
        raise ValueError('PDFとして読み取れません。')
    directory.mkdir(parents=True,exist_ok=True)
    (directory/aid).write_bytes(data)
    return attachment,assets


def add_pdf_reference(doc,data,assets_dir):
    exe=shutil.which('pdftoppm')
    if not exe:
        doc['warnings'].append('PDF比較画像を生成できませんでした。pdftoppmが見つかりません。')
        return
    try:
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if len(reader.pages)!=len(doc['pages']):
            doc['warnings'].append('PDFとIDMLのページ数が異なるため、比較画像の自動対応付けを行いませんでした。')
            return
        with tempfile.TemporaryDirectory(dir=Path(assets_dir).parent) as tmp:
            source=Path(tmp)/'reference.pdf';source.write_bytes(data)
            prefix=Path(tmp)/'page'
            cmd=[exe,'-r','110','-f','1','-l',str(min(100,len(doc['pages']))),'-png',str(source),str(prefix)]
            subprocess.run(cmd,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE,timeout=90,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            rendered=sorted(Path(tmp).glob('page-*.png'),key=lambda p:int(p.stem.split('-')[-1]))
            for page,pdf_page,png in zip(doc['pages'],reader.pages,rendered):
                aid=save_image(png.read_bytes(),assets_dir)
                doc['assets'][aid]={'id':aid,'name':'原版PDF '+page['label'],'kind':'reference'}
                page['reference_asset']=aid
                labels=[]
                def visitor(content,cm,tm,font,size):
                    value=content.strip()
                    if re.fullmatch(r'\d{1,4}',value) and 0<=tm[5]<40:
                        labels.append(value)
                pdf_page.extract_text(visitor_text=visitor)
                if len(set(labels))==1:
                    page['label']=labels[0]
                    page['printed_label']=labels[0]
    except (OSError,ValueError,subprocess.SubprocessError) as e:
        doc['warnings'].append('PDF比較画像の生成に失敗しました。編集用HTMLは利用できます。')
